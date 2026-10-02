"""Genera exceles DESORDENADOS con datos 100% falsos, imitando el caos típico de oficinas de gobierno.

Problemas que se inyectan a propósito (el ETL tiene que resolverlos):
  - encabezados con nombres distintos y en filas distintas
  - fechas como fecha real, texto dd/mm/aaaa, '12-abr-2026', número serial de Excel o en 3 columnas
  - mismo trámite escrito de 5 formas distintas, con typos
  - nombres de encargados con mayúsculas, sin acentos o abreviados
  - folios con formatos distintos
  - folios duplicados entre archivos y dentro del mismo archivo
  - filas vacías, fila de TOTAL, fechas imposibles, celdas en blanco

Uso:  python -m src.generar_exceles
"""
import datetime as dt
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "data" / "raw"

MESES_ABR = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]

TRAMITES = {
    "Acta de nacimiento": ["Acta de nacimiento", "ACTA DE NACIMIENTO", "acta nacimiento",
                           "Acta de Nacimiento ", "Acta nacimiento."],
    "Constancia de residencia": ["Constancia de residencia", "CONSTANCIA RESIDENCIA",
                                 "Const. de residencia", "constancia de residencia "],
    "Licencia de funcionamiento": ["Licencia de funcionamiento", "LICENCIA FUNCIONAMIENTO",
                                   "Licencia func.", "licencia de funcionamiento"],
    "Pago de predial": ["Pago de predial", "PAGO PREDIAL", "Predial", "pago de predial "],
    "Permiso de construcción": ["Permiso de construcción", "Permiso de construccion",
                                "PERMISO CONSTRUCCION", "permiso construcción "],
    "Registro de matrimonio": ["Registro de matrimonio", "REGISTRO MATRIMONIO",
                               "Reg. matrimonio", "registro de matrimonio"],
}
TYPOS = ["Acta de nacimeinto", "Constancia de residensia", "Pago de predal", "Permiso de contruccion"]

ENCARGADOS = {
    "Ana López Ruiz": ["Ana López Ruiz", "ANA LOPEZ RUIZ", "A. López", "Ana Lopez"],
    "Carlos Mendoza Vega": ["Carlos Mendoza Vega", "CARLOS MENDOZA", "C. Mendoza", "carlos mendoza vega"],
    "Lucía Hernández Soto": ["Lucía Hernández Soto", "LUCIA HERNANDEZ", "L. Hernández", "Lucia Hernandez S."],
    "Jorge Ramírez Cruz": ["Jorge Ramírez Cruz", "JORGE RAMIREZ", "J. Ramírez", "Jorge Ramirez"],
    "Marisol Torres Díaz": ["Marisol Torres Díaz", "MARISOL TORRES", "M. Torres", "marisol torres diaz"],
    "Roberto Salinas Peña": ["Roberto Salinas Peña", "ROBERTO SALINAS", "R. Salinas", "Roberto Salinas"],
}

ESTATUS = {
    "Concluido": ["Concluido", "CONCLUIDO", "Concluído", "Finalizado", "Terminado"],
    "En proceso": ["En proceso", "EN PROCESO", "Proceso", "en proceso "],
    "Pendiente": ["Pendiente", "PENDIENTE", "por revisar"],
    "Rechazado": ["Rechazado", "RECHAZADO", "Cancelado"],
}
ESTATUS_PESOS = [("Concluido", 70), ("En proceso", 15), ("Pendiente", 9), ("Rechazado", 6)]

NOMBRES = ["María", "José", "Juan", "Guadalupe", "Francisco", "Luis", "Rosa", "Miguel", "Patricia",
           "Alejandro", "Fernanda", "Ricardo", "Elena", "Daniel", "Sofía", "Héctor", "Paola", "Raúl"]
APELLIDOS = ["García", "Martínez", "Rodríguez", "Pérez", "Sánchez", "Ramírez", "Flores", "Gómez",
             "Morales", "Vázquez", "Jiménez", "Reyes", "Cruz", "Ortiz", "Gutiérrez", "Chávez"]


def _pesos_tramite(fecha: dt.date):
    """Predial pesa más en ene-mar (descuento por pronto pago)."""
    pesos = {"Acta de nacimiento": 25, "Constancia de residencia": 20, "Licencia de funcionamiento": 8,
             "Pago de predial": 30 if fecha.month <= 3 else 12,
             "Permiso de construcción": 7, "Registro de matrimonio": 6}
    return list(pesos), list(pesos.values())


def _dias_habiles(inicio: dt.date, fin: dt.date):
    d = inicio
    while d <= fin:
        if d.weekday() < 5:
            yield d
        d += dt.timedelta(days=1)


def _generar_registros():
    """Registros 'verdaderos' (antes de ensuciarlos). Cada uno: dict con oficina y campos limpios."""
    rnd = random
    ciudadanos = [f"{rnd.choice(NOMBRES)} {rnd.choice(APELLIDOS)} {rnd.choice(APELLIDOS)}" for _ in range(700)]
    registros, folio = [], 0
    for oficina in ("Centro", "Norte", "Sur"):
        for dia in _dias_habiles(dt.date(2026, 1, 5), dt.date(2026, 6, 30)):
            n = rnd.randint(2, 6)
            # Pico real de predial: último día del descuento, solo en Centro
            if oficina == "Centro" and dia == dt.date(2026, 2, 27):
                n += 60
            for _ in range(n):
                tipos, pesos = _pesos_tramite(dia)
                folio += 1
                registros.append({
                    "folio": folio,
                    "oficina": oficina,
                    "fecha": dia,
                    "tramite": "Pago de predial" if (oficina == "Centro" and dia == dt.date(2026, 2, 27) and rnd.random() < .9)
                               else rnd.choices(tipos, pesos)[0],
                    "encargado": rnd.choice(list(ENCARGADOS)),
                    "ciudadano": rnd.choice(ciudadanos),
                    "estatus": rnd.choices([e for e, _ in ESTATUS_PESOS], [p for _, p in ESTATUS_PESOS])[0],
                })
    rnd.shuffle(registros)
    registros.sort(key=lambda r: (r["fecha"], r["folio"]))
    return registros


# ---- ensuciadores -----------------------------------------------------------
def _folio_sucio(n):
    return random.choice([f"TR-2026-{n:05d}", f"TR2026-{n}", f"tr-2026-{n} ", f"TR-2026-{n:05d}"])


def _tramite_sucio(canon):
    if random.random() < 0.03:
        base = [t for t in TYPOS if canon.split()[0].lower()[:3] in t.lower()]
        if base:
            return base[0]
    return random.choice(TRAMITES[canon])


def _encargado_sucio(canon):
    if random.random() < 0.04:
        return None  # sin asignar
    return random.choice(ENCARGADOS[canon])


def _estatus_sucio(canon):
    if random.random() < 0.02:
        return None
    return random.choice(ESTATUS[canon])


def _fila_base(r):
    return [_folio_sucio(r["folio"]), _tramite_sucio(r["tramite"]), _encargado_sucio(r["encargado"]),
            r["ciudadano"], _estatus_sucio(r["estatus"])]


def _fecha_texto_barra(d):
    return d.strftime("%d/%m/%Y")


def _fecha_texto_abr(d):
    return f"{d.day:02d}-{MESES_ABR[d.month - 1]}-{d.year}"


def _serial_excel(d):
    return (d - dt.date(1899, 12, 30)).days


def _encabezado(ws, fila, textos):
    for i, t in enumerate(textos, start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True)


# ---- archivos ---------------------------------------------------------------
def generar(salida: Path = SALIDA, semilla: int = 42):
    random.seed(semilla)
    salida.mkdir(parents=True, exist_ok=True)
    for viejo in salida.glob("*.xlsx"):
        viejo.unlink()

    regs = _generar_registros()
    centro_1 = [r for r in regs if r["oficina"] == "Centro" and r["fecha"].month <= 3]
    centro_2 = [r for r in regs if r["oficina"] == "Centro" and r["fecha"].month > 3]
    norte = [r for r in regs if r["oficina"] == "Norte"]
    sur = [r for r in regs if r["oficina"] == "Sur"]

    # 1) Centro ene-mar: bien armado, fechas reales
    wb = Workbook(); ws = wb.active; ws.title = "Tramites"
    _encabezado(ws, 1, ["Folio", "Fecha", "Trámite", "Encargado", "Ciudadano", "Estatus"])
    for r in centro_1:
        f = _fila_base(r)
        ws.append([f[0], r["fecha"], f[1], f[2], f[3], f[4]])
    # un par de errores de captura en fecha
    ws.append(["TR-2026-99001", dt.date(2062, 3, 4), "Predial", "Ana López Ruiz", "Rosa Cruz Gómez", "Concluido"])
    ws.append(["TR-2026-99002", None, "Acta de nacimiento", "C. Mendoza", "Luis Pérez Reyes", "Concluido"])
    wb.save(salida / "Tramites_Ene-Mar.xlsx")

    # 2) Norte: título arriba, encabezado en fila 4, fechas texto dd/mm/aaaa, filas vacías y TOTAL
    wb = Workbook(); ws = wb.active; ws.title = "Hoja1"
    ws["A1"] = "OFICINA NORTE - REGISTRO DE TRÁMITES 2026"; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Elaboró: Ventanilla única"
    _encabezado(ws, 4, ["No. Folio", "Fecha de trámite", "Tipo de tramite", "Responsable",
                        "Nombre del solicitante", "Estado"])
    for i, r in enumerate(norte):
        f = _fila_base(r)
        ws.append([f[0], _fecha_texto_barra(r["fecha"]), f[1], f[2], f[3], f[4]])
        if i % 90 == 89:
            ws.append([None] * 6)  # fila vacía en medio
    ws.append(["TOTAL", len(norte)])
    wb.save(salida / "registro_oficina_norte.xlsx")

    # 3) Centro abr-jun: dos hojas, MAYÚSCULAS, fechas texto '12-abr-2026' o serial; duplicados de ene-mar
    wb = Workbook(); abril = wb.active; abril.title = "Abril"; resto = wb.create_sheet("Mayo-Junio")
    for hoja in (abril, resto):
        _encabezado(hoja, 1, ["FOLIO", "FECHA", "TRAMITE", "ATENDIO", "SOLICITANTE", "STATUS"])
    for r in centro_2:
        f = _fila_base(r)
        fecha = _serial_excel(r["fecha"]) if random.random() < .25 else _fecha_texto_abr(r["fecha"])
        (abril if r["fecha"].month == 4 else resto).append([f[0], fecha, f[1], f[2], f[3], f[4]])
    # 'reimpresiones': mismos folios que ya estaban en el archivo de ene-mar
    for r in random.sample(centro_1, 25):
        f = _fila_base(r)
        abril.append([f[0], _fecha_texto_abr(r["fecha"]), f[1], f[2], f[3], f[4]])
    wb.save(salida / "TRAMITES 2do trim FINAL (2).xlsx")

    # 4) Sur: fecha partida en Día / Mes / Año, y duplicados exactos dentro del mismo archivo
    wb = Workbook(); ws = wb.active; ws.title = "Datos"
    _encabezado(ws, 1, ["Folio", "Dia", "Mes", "Año", "Tramite", "Atendió", "Solicitante", "Resultado"])
    filas = []
    for r in sur:
        f = _fila_base(r)
        filas.append([f[0], r["fecha"].day, r["fecha"].month, r["fecha"].year, f[1], f[2], f[3], f[4]])
    filas += random.sample(filas, 15)
    for fila in filas:
        ws.append(fila)
    wb.save(salida / "Oficina_Sur_ventanilla.xlsx")

    return sorted(p.name for p in salida.glob("*.xlsx")), len(regs)


if __name__ == "__main__":
    archivos, total = generar()
    print(f"{total} registros 'verdaderos' repartidos en {len(archivos)} archivos sucios:")
    for a in archivos:
        print("  -", a)
